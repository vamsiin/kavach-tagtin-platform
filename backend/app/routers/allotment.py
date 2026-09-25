"""RFID Tag / TIN allotment charts (BRD BR-001, FR-004, "tag / TIN allocation
management"). A chart is division-level reference data issued by RDSO / the
Railway; stations and their adjacent block sections are linked to its rows by
station ID. Charts are never edited -- a revised chart is loaded as a new one and
supersedes the current chart of the same division."""
from __future__ import annotations

import csv
import hashlib
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from .. import models, schemas
from ..allotment.ranges import TAG_MAX, TIN_MAX, RangeError, compact, id_key, parse_ranges
from ..allotment.service import current_rows, find_row
from ..auth.deps import get_current_user
from ..auth.permissions import require
from ..database import get_db
from ..documents.service import file_extension, max_upload_bytes, safe_basename, upload_dir

router = APIRouter(tags=["allotment"])

NEAR_ROWS = 2   # rows this close on the chart are neighbouring sections; their numbers must not overlap
CHART_EXT = (".pdf", ".xlsx", ".xls", ".csv", ".png", ".jpg", ".jpeg")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# Row validation
# ---------------------------------------------------------------------------
def _check_rows(rows: list[schemas.AllotmentRowIn]) -> tuple[list[schemas.AllotmentRowOut], list[str], list[str]]:
    """Returns (rows with counts/problems, errors, warnings). Errors block saving;
    warnings are reported (a chart is an issued document -- a total that doesn't add
    up is recorded as printed and flagged, not silently corrected)."""
    out, errors, warnings = [], [], []
    parsed: list[tuple[str, set[int], set[int]]] = []
    seen_ids: dict[int, str] = {}
    for i, r in enumerate(rows, 1):
        label = f"Row {r.s_no or i} ({r.name or '?'})"
        probs: list[str] = []
        if not (r.name or "").strip():
            errors.append(f"Row {r.s_no or i}: station / block section name is required")
        tags, tins = [], []
        try:
            tags = parse_ranges(r.tag_ranges, TAG_MAX)
        except RangeError as e:
            errors.append(f"{label}: tag range {e}")
        try:
            tins = parse_ranges(r.tin_ranges, TIN_MAX)
        except RangeError as e:
            errors.append(f"{label}: TIN range {e}")
        if r.tag_total is not None and tags and r.tag_total != len(tags):
            probs.append(f"Total printed as {r.tag_total}, but the tag ranges give {len(tags)}")
        k = id_key(r.station_id)
        if k is not None:
            if k in seen_ids:
                errors.append(f"{label}: station ID {r.station_id} is also used by {seen_ids[k]}")
            seen_ids[k] = r.name
        parsed.append((r.name, set(tags), set(tins)))
        warnings.extend(f"{label}: {p}" for p in probs)
        out.append(schemas.AllotmentRowOut(**r.model_dump(), tag_count=len(tags), tin_count=len(tins),
                                           problems=probs))
    # Numbers are deliberately reused across a division (a tag ID may repeat beyond
    # 25 km, H2.14), so only overlaps between rows close together on the chart --
    # neighbouring stations and block sections, listed in line order -- are flagged.
    for i, (name_a, tags_a, tins_a) in enumerate(parsed):
        for name_b, tags_b, tins_b in parsed[i + 1:i + 1 + NEAR_ROWS]:
            if tags_a & tags_b:
                warnings.append(f"Tag numbers {compact(tags_a & tags_b)} are in both {name_a} and {name_b}, "
                                f"which are close together on the chart")
            if tins_a & tins_b:
                warnings.append(f"TIN {compact(tins_a & tins_b)} is in both {name_a} and {name_b}, "
                                f"which are close together on the chart")
    return out, errors, warnings


def _chart_out(c: models.AllotmentChartORM) -> schemas.AllotmentChartOut:
    return schemas.AllotmentChartOut(
        id=c.id, title=c.title, division=c.division, reference_no=c.reference_no, issued_on=c.issued_on,
        issued_by=c.issued_by, remarks=c.remarks, is_current=bool(c.is_current),
        source_filename=c.source_filename, source_sha256=c.source_sha256, created_by=c.created_by,
        created_at=c.created_at, superseded_by_id=c.superseded_by_id, row_count=len(c.rows))


def _row_in(r: models.AllotmentRowORM) -> schemas.AllotmentRowIn:
    return schemas.AllotmentRowIn(**{f: getattr(r, f) for f in schemas.AllotmentRowIn.model_fields})


def _get_chart(chart_id: int, db: Session) -> models.AllotmentChartORM:
    c = db.get(models.AllotmentChartORM, chart_id)
    if not c:
        raise HTTPException(404, "Allotment chart not found")
    return c


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------
@router.get("/allotment-charts", response_model=list[schemas.AllotmentChartOut])
def list_charts(db: Session = Depends(get_db), user: models.UserORM = Depends(get_current_user)):
    return [_chart_out(c) for c in db.query(models.AllotmentChartORM).order_by(models.AllotmentChartORM.id.desc())]


@router.get("/allotment-charts/{chart_id}", response_model=schemas.AllotmentChartDetail)
def get_chart(chart_id: int, db: Session = Depends(get_db), user: models.UserORM = Depends(get_current_user)):
    c = _get_chart(chart_id, db)
    rows, _, warnings = _check_rows([_row_in(r) for r in c.rows])
    for out, r in zip(rows, c.rows):
        out.id = r.id
    return schemas.AllotmentChartDetail(**_chart_out(c).model_dump(), rows=rows, warnings=warnings)


@router.post("/allotment-charts", response_model=schemas.AllotmentChartDetail)
def create_chart(payload: schemas.AllotmentChartCreate, db: Session = Depends(get_db),
                 user: models.UserORM = Depends(get_current_user)):
    require(user, "allotment.manage", "Loading an allotment chart")
    if not payload.title.strip() or not payload.division.strip():
        raise HTTPException(422, "Title and division are required")
    if not payload.rows:
        raise HTTPException(422, "The chart has no rows")
    _, errors, _ = _check_rows(payload.rows)
    if errors:
        raise HTTPException(422, {"message": "The chart has errors", "errors": errors})
    c = models.AllotmentChartORM(
        title=payload.title.strip(), division=payload.division.strip(), reference_no=payload.reference_no,
        issued_on=payload.issued_on, issued_by=payload.issued_by, remarks=payload.remarks,
        is_current=payload.supersede_current, created_by=user.username, created_at=_now())
    c.rows = [models.AllotmentRowORM(**{**r.model_dump(), "name": r.name.strip()}) for r in payload.rows]
    db.add(c)
    db.flush()
    if payload.supersede_current:
        _supersede(db, c)
    db.commit()
    db.refresh(c)
    return get_chart(c.id, db, user)


def _supersede(db: Session, c: models.AllotmentChartORM) -> None:
    for old in (db.query(models.AllotmentChartORM)
                .filter(models.AllotmentChartORM.id != c.id, models.AllotmentChartORM.is_current.is_(True))):
        if old.division.strip().lower() == c.division.strip().lower():
            old.is_current = False
            old.superseded_by_id = c.id


@router.post("/allotment-charts/{chart_id}/make-current", response_model=schemas.AllotmentChartOut)
def make_current(chart_id: int, db: Session = Depends(get_db), user: models.UserORM = Depends(get_current_user)):
    """Roll back to an earlier chart, e.g. if a revision was loaded by mistake."""
    require(user, "allotment.manage", "Changing the current allotment chart")
    c = _get_chart(chart_id, db)
    c.is_current, c.superseded_by_id = True, None
    _supersede(db, c)
    db.commit()
    db.refresh(c)
    return _chart_out(c)


@router.post("/allotment-charts/{chart_id}/file", response_model=schemas.AllotmentChartOut)
async def attach_file(chart_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                      user: models.UserORM = Depends(get_current_user)):
    """Keep the issued chart (usually the signed RDSO letter) with the transcribed rows."""
    require(user, "allotment.manage", "Attaching an allotment chart file")
    c = _get_chart(chart_id, db)
    name = safe_basename(file.filename)
    ext = file_extension(name)
    if ext not in CHART_EXT:
        raise HTTPException(422, f"File type {ext or '(none)'} not accepted; use one of {', '.join(CHART_EXT)}")
    data = await file.read()
    if not data:
        raise HTTPException(422, "The file is empty")
    if len(data) > max_upload_bytes():
        raise HTTPException(413, "The file is too large")
    rel = Path("allotment") / f"chart_{c.id}" / f"{uuid.uuid4().hex}{ext}"
    path = upload_dir() / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    c.source_filename, c.source_path = name, rel.as_posix()
    c.source_sha256 = hashlib.sha256(data).hexdigest()
    db.commit()
    db.refresh(c)
    return _chart_out(c)


@router.get("/allotment-charts/{chart_id}/file")
def download_file(chart_id: int, db: Session = Depends(get_db), user: models.UserORM = Depends(get_current_user)):
    c = _get_chart(chart_id, db)
    if not c.source_path:
        raise HTTPException(404, "No file attached to this chart")
    p = (upload_dir() / c.source_path).resolve()
    if upload_dir() not in p.parents or not p.exists():
        raise HTTPException(404, "Chart file not found")
    return FileResponse(p, filename=c.source_filename or p.name)


# ---------------------------------------------------------------------------
# Import from a spreadsheet (preview only -- nothing is saved)
# ---------------------------------------------------------------------------
_COLS = [  # (field, header keywords -- all must appear), most specific first
    ("s_no", ("s", "no")), ("division_code", ("division",)), ("station_code", ("station", "code")),
    ("station_id", ("station", "id")), ("tin_allotted", ("allot",)), ("tin_ranges", ("tin",)),
    ("tag_total", ("total",)), ("tag_ranges", ("tag",)), ("name", ("station",)), ("name", ("section",)),
]


def _map_header(cells: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, raw in enumerate(cells):
        h = " ".join(str(raw or "").lower().replace(".", " ").replace("/", " ").split())
        words = set(h.split())
        for field, keys in _COLS:
            if field in found:
                continue
            if all(k in words or (len(k) > 2 and k in h) for k in keys):
                found[field] = idx
                break
    return found


def _read_table(name: str, data: bytes) -> list[list[str]]:
    ext = file_extension(name)
    if ext == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    raise HTTPException(422, "Import a .xlsx or .csv file (a scanned PDF chart has to be transcribed first; "
                             "attach the PDF to the chart afterwards)")


@router.post("/allotment-charts/parse", response_model=schemas.AllotmentParseOut)
async def parse_chart(file: UploadFile = File(...), user: models.UserORM = Depends(get_current_user)):
    require(user, "allotment.manage", "Importing an allotment chart")
    data = await file.read()
    table = _read_table(safe_basename(file.filename), data)
    hdr_idx, cols = None, {}
    for i, row in enumerate(table[:15]):
        cols = _map_header(row)
        if "name" in cols and "tag_ranges" in cols:
            hdr_idx = i
            break
    if hdr_idx is None:
        raise HTTPException(422, "Couldn't find the header row: it needs a 'Station / Block Section' "
                                 "column and an 'RFID Tag ... Range' column")
    rows: list[schemas.AllotmentRowIn] = []
    for row in table[hdr_idx + 1:]:
        get = lambda f: (row[cols[f]].strip() if f in cols and cols[f] < len(row) and row[cols[f]] is not None else "")
        if not get("name"):
            continue
        total = get("tag_total")
        rows.append(schemas.AllotmentRowIn(
            s_no=get("s_no") or None, name=get("name"), division_code=get("division_code") or None,
            station_code=get("station_code") or None, station_id=get("station_id") or None,
            tin_ranges=get("tin_ranges") or None, tin_allotted=get("tin_allotted") or None,
            tag_ranges=get("tag_ranges") or None,
            tag_total=int(float(total)) if total.replace(".", "", 1).isdigit() else None))
    checked, errors, warnings = _check_rows(rows)
    return schemas.AllotmentParseOut(rows=checked, errors=errors, warnings=warnings)


# ---------------------------------------------------------------------------
# A station's allotment and usage
# ---------------------------------------------------------------------------
def _row_out(r: models.AllotmentRowORM | None) -> schemas.AllotmentRowOut | None:
    if r is None:
        return None
    out, _, _ = _check_rows([_row_in(r)])
    out[0].id = r.id
    return out[0]


def _usage(own: set[int], adjacent: list[tuple[str, set[int]]], used: list[tuple[int, str]]) -> schemas.AllotmentUsage:
    used_nums = {n for n, _ in used}
    free = sorted(own - used_nums)

    def owner(n):
        if n in own:
            return "own"
        return next((name for name, s in adjacent if n in s), None)
    uses = [schemas.NumberUse(number=n, ref=ref, owner=owner(n)) for n, ref in sorted(used)]
    return schemas.AllotmentUsage(
        allotted=len(own), used_in_own=len(own & used_nums), free=len(free), free_ranges=compact(free),
        next_free_odd=next((n for n in free if n % 2), None), next_free_even=next((n for n in free if not n % 2), None),
        used=uses, outside=[u for u in uses if u.owner is None])


@router.get("/stations/{station_id}/allotment", response_model=schemas.StationAllotmentOut)
def station_allotment(station_id: int, db: Session = Depends(get_db),
                      user: models.UserORM = Depends(get_current_user)):
    st = db.get(models.StationORM, station_id)
    if not st:
        raise HTTPException(404, "Station not found")
    rows = current_rows(db)
    if not rows:
        return schemas.StationAllotmentOut()
    own = find_row(rows, st.kavach_station_id, st.name)
    adj = [(a.name, find_row(rows, a.kavach_station_id, a.name)) for a in st.adjacent_stations]

    def nums(r, attr, hi):
        try:
            return set(parse_ranges(getattr(r, attr), hi)) if r is not None else set()
        except RangeError:
            return set()
    out = schemas.StationAllotmentOut(
        chart=_chart_out(own.chart if own is not None else rows[0].chart), own=_row_out(own),
        adjacent=[schemas.AdjacentAllotment(adjacent_name=n, row=_row_out(r)) for n, r in adj])
    if own is not None:
        out.tags = _usage(nums(own, "tag_ranges", TAG_MAX),
                          [(r.name, nums(r, "tag_ranges", TAG_MAX)) for _, r in adj if r is not None],
                          [(t.tag_number, t.tag_ref) for t in st.tags if t.tag_number is not None])
        out.tins = _usage(nums(own, "tin_ranges", TIN_MAX),
                          [(r.name, nums(r, "tin_ranges", TIN_MAX)) for _, r in adj if r is not None],
                          [(n.tin_no, f"TIN {n.tin_no}") for n in st.tins])
    return out
