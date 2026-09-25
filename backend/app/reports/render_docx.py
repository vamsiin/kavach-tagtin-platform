"""Word (.docx) renderer for app.reports.model.Report (python-docx, A4)."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from . import model as m

NAVY = RGBColor(0x1F, 0x38, 0x64)
MUTED = RGBColor(0x6B, 0x72, 0x80)
TONE_BG = {"ok": "EAF7EE", "fail": "FDECEA", "warn": "FDF6E3", "info": "EEF2F7"}
TONE_FG = {"ok": RGBColor(0x1E, 0x7D, 0x3C), "fail": RGBColor(0xC0, 0x39, 0x2B),
           "warn": RGBColor(0x8A, 0x61, 0x16), "info": NAVY}


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _cell_text(cell, text: str, bold=False, color: RGBColor | None = None, size=8.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = str(text if text is not None else "").split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if i < len(lines) - 1:
            run.add_break()


def _page_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", instr), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = attr
        else:
            el.set(qn("w:fldCharType"), attr)
        run._r.append(el)
    run.font.size = Pt(7)


def render_docx(report: m.Report, wide: bool = False) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    if wide:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(15))
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(16)
    avail_mm = (297 if wide else 210) - 30

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)

    # Header / footer
    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"KAVACH RFID Tag/TIN Layout  |  {report.title}  |  "
                    f"{report.station_name} ({report.station_code})  |  Layout {report.layout_version}")
    hr.font.size, hr.font.color.rgb = Pt(8), NAVY
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"{report.report_id}  |  Generated {report.generated_at[:16].replace('T', ' ')} UTC "
                    f"by {report.generated_by}  |  Uncontrolled when printed  |  Page ")
    fr.font.size, fr.font.color.rgb = Pt(7), MUTED
    _page_field(fp, "PAGE")

    t = doc.add_paragraph()
    tr = t.add_run(report.title)
    tr.bold, tr.font.size, tr.font.color.rgb = True, Pt(18), NAVY
    _kv(doc, [("Station", f"{report.station_name} (KAVACH ID {report.station_code})"),
              ("Layout version", report.layout_version),
              ("Workflow status", report.station_status),
              ("Report ID", report.report_id)], avail_mm)

    for b in report.blocks:
        _block(doc, b, avail_mm)

    buf = io.BytesIO()
    doc.core_properties.title = report.title
    doc.core_properties.author = report.generated_by
    doc.core_properties.subject = report.report_id
    doc.save(buf)
    return buf.getvalue()


def _set_widths(table, widths_mm):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_mm):
            row.cells[i].width = Mm(w)


def _kv(doc, rows, avail_mm):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for k, v in rows:
        c = tbl.add_row().cells
        _cell_text(c[0], k, bold=True)
        _shade(c[0], "F5F6F8")
        _cell_text(c[1], v)
    _set_widths(tbl, [avail_mm * 0.28, avail_mm * 0.72])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _block(doc, b, avail_mm):
    if isinstance(b, m.Heading):
        p = doc.add_paragraph()
        r = p.add_run(b.text)
        r.bold, r.font.color.rgb = True, NAVY
        r.font.size = Pt(13 if b.level == 1 else 11)
        p.paragraph_format.space_before = Pt(10 if b.level == 1 else 6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
    elif isinstance(b, m.Para):
        p = doc.add_paragraph()
        r = p.add_run(b.text)
        if b.style in ("muted", "small"):
            r.font.color.rgb = MUTED
        if b.style == "small":
            r.font.size = Pt(8)
    elif isinstance(b, m.Banner):
        tbl = doc.add_table(rows=1, cols=1)
        c = tbl.rows[0].cells[0]
        _cell_text(c, b.text, bold=True, color=TONE_FG[b.tone], size=9.5)
        _shade(c, TONE_BG[b.tone])
        _set_widths(tbl, [avail_mm])
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif isinstance(b, m.KeyValues):
        _kv(doc, b.rows, avail_mm)
    elif isinstance(b, m.Bullets):
        for item in b.items:
            doc.add_paragraph(item, style="List Bullet")
    elif isinstance(b, m.PageBreak):
        doc.add_page_break()
    elif isinstance(b, m.SignOff):
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        for i, h in enumerate(("Role", "Name", "Date", "Signature")):
            c = tbl.rows[0].cells[i]
            _cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade(c, "1F3864")
        for r_, n, d in b.rows:
            cells = tbl.add_row().cells
            for i, v in enumerate((r_, n, d, "")):
                _cell_text(cells[i], v)
            tbl.rows[-1].height = Mm(12)
        _set_widths(tbl, [avail_mm * w for w in (0.22, 0.28, 0.2, 0.3)])
        doc.add_paragraph()
    elif isinstance(b, m.Table):
        if not b.rows:
            p = doc.add_paragraph()
            r = p.add_run(b.empty_text)
            r.font.color.rgb = MUTED
            return
        widths = b.widths or [1] * len(b.columns)
        total = sum(widths)
        tbl = doc.add_table(rows=1, cols=len(b.columns))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(b.columns):
            c = tbl.rows[0].cells[i]
            _cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade(c, "1F3864")
        _repeat_header(tbl.rows[0])
        for ri, row in enumerate(b.rows):
            cells = tbl.add_row().cells
            for ci, v in enumerate(row):
                tone = m.tone_for(v) if ci == b.status_col else None
                _cell_text(cells[ci], v, bold=bool(tone), color=TONE_FG[tone] if tone else None)
                if tone:
                    _shade(cells[ci], TONE_BG[tone])
                elif ri % 2 == 1:
                    _shade(cells[ci], "F5F6F8")
        _set_widths(tbl, [avail_mm * w / total for w in widths])
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
