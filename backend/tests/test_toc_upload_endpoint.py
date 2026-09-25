"""
Proves the actual HTTP upload endpoint (POST /stations/{id}/verify-toc) works for
each supported format, not just the underlying parser functions in isolation
(test_toc_multiformat.py) -- multipart file upload handling, extension dispatch,
and the .doc rejection path all live in routers/verify.py, not in parser.py.
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import io
import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("JWT_SECRET", "test-secret-for-toc-upload-tests")


@pytest.fixture
def client(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path}")

    for mod in list(sys.modules):
        if mod == "app" or mod.startswith("app."):
            del sys.modules[mod]

    from app.main import app as fastapi_app
    with TestClient(fastapi_app) as c:
        yield c


def _auth_headers(client):
    client.post("/auth/register", json={
        "username": "admin", "full_name": "Admin", "role": "admin", "password": "adminpass123",
    })
    r = client.post("/auth/login", json={"username": "admin", "password": "adminpass123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _make_station_with_tag(client, headers) -> int:
    sid = client.post("/stations", json={"name": "T", "kavach_station_id": "1"},
                       headers=headers).json()["id"]
    client.post(f"/stations/{sid}/tags",
                json={"tag_ref": "R-775", "tag_type": "S", "location_m": 262750},
                headers=headers)
    client.post(f"/stations/{sid}/tags",
                json={"tag_ref": "R-777", "tag_type": "N", "location_m": 263569},
                headers=headers)
    return sid


ROW_1 = [
    "1", "S1D", "S1 ID", "Station", "DN", "", "", "", "", "", "", "", "", "",
    "R-775", "(0#R-775#0)#(819#R-777#0)", "819", "2", "", "", "",
]


def test_upload_csv(client):
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)

    csv_content = ",".join([""] * 21) + "\n" + ",".join([""] * 21) + "\n" + ",".join(ROW_1) + "\n"
    files = {"toc_file": ("toc.csv", io.BytesIO(csv_content.encode()), "text/csv")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files, headers=headers)
    assert r.status_code == 200
    assert r.json()["summary"]["total_checks"] > 0


def test_upload_xlsx(client):
    import openpyxl
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([""] * 21)
    ws.append(ROW_1)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    files = {"toc_file": ("toc.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files, headers=headers)
    assert r.status_code == 200
    assert r.json()["summary"]["total_checks"] > 0


def test_upload_docx(client):
    import docx
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)

    doc = docx.Document()
    table = doc.add_table(rows=0, cols=21)
    for row_data in [[""] * 21, ROW_1]:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = val
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    files = {"toc_file": ("toc.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files, headers=headers)
    assert r.status_code == 200
    assert r.json()["summary"]["total_checks"] > 0


def test_upload_xls(client):
    import xlwt
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)

    wb = xlwt.Workbook()
    ws = wb.add_sheet("ToC")
    for r_idx, row in enumerate([[""] * 21, ROW_1]):
        for c_idx, val in enumerate(row):
            ws.write(r_idx, c_idx, val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    files = {"toc_file": ("toc.xls", buf, "application/vnd.ms-excel")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files, headers=headers)
    assert r.status_code == 200
    assert r.json()["summary"]["total_checks"] > 0


def test_upload_legacy_doc_rejected_with_actionable_message(client):
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)

    files = {"toc_file": ("toc.doc", io.BytesIO(b"fake legacy doc bytes"), "application/msword")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files, headers=headers)
    assert r.status_code == 422
    assert ".docx" in r.json()["detail"]


def test_upload_unsupported_extension_rejected(client):
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)

    files = {"toc_file": ("toc.pdf", io.BytesIO(b"irrelevant"), "application/pdf")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files, headers=headers)
    assert r.status_code == 422


def test_verify_toc_requires_authentication(client):
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)
    files = {"toc_file": ("toc.csv", io.BytesIO(b"irrelevant"), "text/csv")}
    r = client.post(f"/stations/{sid}/verify-toc", files=files)  # no auth header
    assert r.status_code == 401


def test_verify_requires_authentication(client):
    headers = _auth_headers(client)
    sid = _make_station_with_tag(client, headers)
    r = client.get(f"/stations/{sid}/verify")  # no auth header
    assert r.status_code == 401
