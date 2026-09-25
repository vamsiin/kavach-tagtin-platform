"""
Tests the non-CSV ToC parsers by actually generating real files in each format
(via openpyxl / python-docx -- the same libraries the parsers themselves use to
read) and parsing them back, rather than just trusting the parser code reads
correctly. This proves the round-trip works for content shaped like the real
ALER_ToC.csv, though see parser.py's module docstring for the honest caveat that
none of these have been tested against a genuine Railway/vendor-exported .xlsx or
.docx file, only synthetic ones built for this test.
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.toc.parser import (
    parse_toc_xlsx, parse_toc_docx, parse_toc_file, UnsupportedFormatError,
)

# A couple of real rows from the ALER ToC, reused across formats for a fair
# comparison against the already-proven CSV parsing of the same underlying data.
HEADER_ROW = [""] * 21
SUBHEADER_ROW = [""] * 21
ROW_1 = [
    "1", "S1D", "S1 ID", "Station", "DN", "", "", "", "", "", "", "", "", "",
    "R-775", "(0#R-775#0)#(819#R-777#0)#(200#R-779#0)", "1019", "2", "", "", "",
]
ROW_2 = [
    "2", "S1 ID", "S1", "Station", "DN", "", "", "", "", "", "", "", "", "",
    "R-779", "(0#R-779#0)#(821#R-781#0)#(209#R-783#0)", "1030", "2", "", "", "",
]


def test_parse_xlsx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [HEADER_ROW, SUBHEADER_ROW, ROW_1, ROW_2]:
        ws.append(row)
    path = tmp_path / "toc.xlsx"
    wb.save(path)

    routes = parse_toc_xlsx(str(path))
    assert len(routes) == 2
    assert routes[0].s_no == "1"
    assert routes[0].entry_signal == "S1D"
    assert routes[0].tag_ids_in_order == ["R-775", "R-777", "R-779"]
    assert routes[0].stated_distance_m == 1019
    assert routes[0].computed_distance_m == 1019
    assert routes[1].tag_ids_in_order == ["R-779", "R-781", "R-783"]


def test_parse_xls(tmp_path):
    """xlwt (write-only, legacy) is used here purely to construct a real .xls file
    for the parser to read back -- it is NOT a runtime dependency of the app itself,
    only of this test (the app only ever needs to READ .xls uploads, via xlrd)."""
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("ToC")
    for r_idx, row in enumerate([HEADER_ROW, SUBHEADER_ROW, ROW_1, ROW_2]):
        for c_idx, val in enumerate(row):
            ws.write(r_idx, c_idx, val)
    path = tmp_path / "toc.xls"
    wb.save(str(path))

    from app.toc.parser import parse_toc_xls
    routes = parse_toc_xls(str(path))
    assert len(routes) == 2
    assert routes[0].s_no == "1"
    assert routes[0].tag_ids_in_order == ["R-775", "R-777", "R-779"]
    assert routes[0].computed_distance_m == 1019
    assert routes[1].tag_ids_in_order == ["R-779", "R-781", "R-783"]


def test_parse_xlsx_numeric_cells_not_mangled(tmp_path):
    """openpyxl hands back a genuine int/float for numeric-looking cells rather
    than a string -- confirms the S.No and distance columns still parse correctly
    when Excel has typed them as numbers rather than text (a very likely real-world
    case: someone typing '1' into a spreadsheet cell without formatting it as text)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    row = list(ROW_1)
    row[0] = 1          # int, not "1"
    row[16] = 1019       # int, not "1019"
    for r in [HEADER_ROW, SUBHEADER_ROW, row]:
        ws.append(r)
    path = tmp_path / "toc_numeric.xlsx"
    wb.save(path)

    routes = parse_toc_xlsx(str(path))
    assert len(routes) == 1
    assert routes[0].s_no == "1"  # not "1.0"
    assert routes[0].stated_distance_m == 1019


def test_parse_docx_single_table(tmp_path):
    import docx
    doc = docx.Document()
    table = doc.add_table(rows=0, cols=21)
    for row_data in [HEADER_ROW, SUBHEADER_ROW, ROW_1, ROW_2]:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = val
    path = tmp_path / "toc.docx"
    doc.save(path)

    routes = parse_toc_docx(str(path))
    assert len(routes) == 2
    assert routes[0].tag_ids_in_order == ["R-775", "R-777", "R-779"]
    assert routes[0].computed_distance_m == 1019
    assert routes[1].s_no == "2"


def test_parse_docx_multiple_tables_concatenated(tmp_path):
    """A real Word export might split routes across multiple tables (e.g. one per
    page) -- confirms rows from every table in the document are picked up, not just
    the first."""
    import docx
    doc = docx.Document()
    t1 = doc.add_table(rows=0, cols=21)
    for row_data in [HEADER_ROW, ROW_1]:
        cells = t1.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = val

    doc.add_paragraph("Continued on next table")

    t2 = doc.add_table(rows=0, cols=21)
    for row_data in [HEADER_ROW, ROW_2]:
        cells = t2.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = val

    path = tmp_path / "toc_multi.docx"
    doc.save(path)

    routes = parse_toc_docx(str(path))
    assert len(routes) == 2
    assert {r.s_no for r in routes} == {"1", "2"}


def test_parse_toc_file_dispatches_by_extension(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [HEADER_ROW, ROW_1]:
        ws.append(row)
    path = tmp_path / "toc.xlsx"
    wb.save(path)

    routes = parse_toc_file(str(path), "toc.xlsx")
    assert len(routes) == 1


def test_parse_toc_file_rejects_legacy_doc_with_clear_message(tmp_path):
    """Legacy .doc is deliberately NOT parsed (see parser.py module docstring for
    why) -- must fail loudly and actionably, not silently produce wrong/empty data."""
    fake_doc_path = tmp_path / "toc.doc"
    fake_doc_path.write_bytes(b"not a real doc file, content irrelevant for this test")

    try:
        parse_toc_file(str(fake_doc_path), "toc.doc")
        assert False, "expected UnsupportedFormatError"
    except UnsupportedFormatError as e:
        assert ".docx" in str(e) or "xlsx" in str(e)  # tells the user what to do instead


def test_parse_toc_file_rejects_unknown_extension(tmp_path):
    path = tmp_path / "toc.txt"
    path.write_text("irrelevant")
    try:
        parse_toc_file(str(path), "toc.txt")
        assert False, "expected UnsupportedFormatError"
    except UnsupportedFormatError:
        pass


def test_content_based_row_detection_ignores_header_position():
    """The row-identification rule (S.No must be a plain digit string) should work
    regardless of how many header rows precede the data -- proves this isn't
    secretly still relying on 'skip exactly 2 rows' from the CSV-specific code path
    that used to exist before this refactor."""
    from app.toc.parser import _rows_to_routes
    rows = [
        ["Some", "Random", "Header", "Text"],  # 1 header row, not 2
        ROW_1,
    ]
    routes = _rows_to_routes(rows)
    assert len(routes) == 1
    assert routes[0].s_no == "1"
