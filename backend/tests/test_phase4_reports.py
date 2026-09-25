"""
Phase 4: BRD FR-007 / Section 11 reports (8 types x PDF + Word), generation log,
verification reports archived automatically at workflow steps, and the revision table.
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import io
import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("JWT_SECRET", "test-secret-for-phase4-tests")

ALL_TYPES = ["layout", "verification", "compliance", "observations", "approval", "audit",
             "traceability", "revisions"]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    for mod in list(sys.modules):
        if mod == "app" or mod.startswith("app."):
            del sys.modules[mod]
    from app.main import app as fastapi_app
    with TestClient(fastapi_app) as c:
        c.upload_root = tmp_path / "uploads"
        yield c


@pytest.fixture
def h(client):
    client.post("/auth/register", json={"username": "admin", "full_name": "Admin",
                                        "role": "admin", "password": "adminpass123"})

    def login(u, p):
        return {"Authorization": "Bearer " + client.post(
            "/auth/login", json={"username": u, "password": p}).json()["access_token"]}
    admin = login("admin", "adminpass123")
    for u, role in (("dana", "designer"), ("cara", "checker"), ("ana", "approver")):
        client.post("/auth/users", json={"username": u, "full_name": u.title(), "role": role,
                                         "password": "pass12345"}, headers=admin)
    return {"admin": admin, **{u: login(u, "pass12345") for u in ("dana", "cara", "ana")}}


@pytest.fixture
def aler(client, h):
    """A realistic Aler-like station (chainage in metres), traced to inputs."""
    d = h["dana"]
    sid = client.post("/stations", json={
        "name": "Aler", "kavach_station_id": "103", "sip_drawing_number": "IBU.1100A",
        "kavach_version": "2", "sip_version": "0", "guideline_version": "3"}, headers=d).json()["id"]
    sip = client.post(f"/stations/{sid}/documents", files={"file": ("sip.pdf", io.BytesIO(b"%PDF"))},
                      data={"doc_type": "SIP"}, headers=d).json()
    e = client.post(f"/stations/{sid}/elements", json={
        "element_ref": "S3", "element_type": "Mainline Starter", "location_m": 262910,
        "source_document_id": sip["id"]}, headers=d).json()
    client.post(f"/stations/{sid}/tags", json={
        "tag_ref": "R-793", "tag_type": "S", "location_m": 262910, "direction": "Nominal",
        "linked_element_id": e["id"], "source_document_id": sip["id"]}, headers=d)
    client.post(f"/stations/{sid}/tins", json={
        "tin_no": 171, "section_type": "Station", "start_m": 262910, "end_m": 263152}, headers=d)
    return sid


def test_report_types_lists_all_eight(client, h):
    r = client.get("/report-types", headers=h["dana"]).json()
    assert [t["code"] for t in r] == ALL_TYPES
    assert all(t["formats"] == ["pdf", "docx"] for t in r)


@pytest.mark.parametrize("code", ALL_TYPES)
def test_every_report_renders_valid_pdf_and_docx(client, h, aler, code):
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
        r = client.get(f"/stations/{aler}/reports/{code}", params={"format": fmt}, headers=h["cara"])
        assert r.status_code == 200, r.text[:300]
        assert r.content[:4].startswith(magic)
        assert f".{fmt}" in r.headers["content-disposition"]
        assert r.headers["x-report-id"].startswith("KVR-103-")


def test_docx_content_matches_data(client, h, aler):
    import docx
    r = client.get(f"/stations/{aler}/reports/layout", params={"format": "docx"}, headers=h["dana"])
    doc = docx.Document(io.BytesIO(r.content))
    text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "R-793" in text and "262910" in text and "IBU.1100A" in text and "SIP r1" in text


def test_report_builders_content(client, h, aler):
    """Content-level checks on the format-neutral model (no PDF parsing)."""
    from app.database import SessionLocal
    from app import models
    from app.routers.reports import build_report
    db = SessionLocal()
    try:
        st = db.get(models.StationORM, aler)
        user = db.query(models.UserORM).filter_by(username="cara").first()
        ver = build_report(db, st, user, "verification").plain_text()
        assert "All" in ver and "automated checks passed" in ver and "H2.6" in ver
        assert "H2.19" in ver  # manual clauses listed
        comp = build_report(db, st, user, "compliance").plain_text()
        assert "FR-001 required inputs uploaded | FAIL" in comp  # survey + chart missing
        assert "MANUAL" in comp
        trace = build_report(db, st, user, "traceability").plain_text()
        assert "R-793" in trace and "SIP rev 1" in trace
    finally:
        db.close()


def test_unknown_type_and_format(client, h, aler):
    assert client.get(f"/stations/{aler}/reports/nope", headers=h["dana"]).status_code == 404
    assert client.get(f"/stations/{aler}/reports/layout", params={"format": "xls"},
                      headers=h["dana"]).status_code == 422
    assert client.get("/stations/999/reports/layout", headers=h["dana"]).status_code == 404
    assert client.get(f"/stations/{aler}/reports/layout").status_code == 401


def test_generation_is_logged(client, h, aler):
    client.get(f"/stations/{aler}/reports/audit", params={"format": "docx"}, headers=h["ana"])
    log = client.get(f"/stations/{aler}/report-log", headers=h["dana"]).json()
    row = log[0]
    assert row["report_type"] == "audit" and row["fmt"] == "docx" and row["trigger"] == "manual"
    assert row["generated_by"].startswith("Ana") and len(row["sha256"]) == 64
    assert row["archived"] is False
    # non-archived reports can't be downloaded from the log
    assert client.get(f"/stations/{aler}/report-log/{row['id']}/download",
                      headers=h["dana"]).status_code == 409


def test_verification_report_archived_at_each_workflow_step(client, h, aler):
    sid = aler
    client.post(f"/stations/{sid}/self-check", json={}, headers=h["dana"])
    client.post(f"/stations/{sid}/submit", json={"assigned_checker_username": "cara"}, headers=h["dana"])
    client.post(f"/stations/{sid}/check", json={}, headers=h["cara"])
    client.post(f"/stations/{sid}/approve", json={}, headers=h["ana"])
    client.post(f"/stations/{sid}/railway-submit",
                json={"reference_no": "R/1", "submitted_on": "2026-09-25"}, headers=h["ana"])
    arch = client.get(f"/stations/{sid}/report-log", params={"archived_only": True},
                      headers=h["dana"]).json()
    assert [a["trigger"] for a in arch] == ["railway_submit", "approve", "check", "self_check"]
    assert all(a["workflow_event_id"] for a in arch)
    wf = client.get(f"/stations/{sid}/workflow", headers=h["dana"]).json()["history"]
    assert {a["workflow_event_id"] for a in arch} == {e["id"] for e in wf if e["action"] != "submit"}
    r = client.get(f"/stations/{sid}/report-log/{arch[-1]['id']}/download", headers=h["dana"])
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    import hashlib
    assert hashlib.sha256(r.content).hexdigest() == arch[-1]["sha256"]
    # layout version now recorded on every workflow event
    assert all(e["layout_version"] == "2.0.3" for e in wf)


def test_archive_failure_never_blocks_workflow(client, h, aler, monkeypatch):
    import app.routers.reports as rep

    def boom(*a, **k):
        raise RuntimeError("renderer exploded")
    monkeypatch.setattr(rep, "render", boom)
    r = client.post(f"/stations/{aler}/self-check", json={}, headers=h["dana"])
    assert r.status_code == 200 and r.json()["status"] == "self_checked"


def test_revision_table(client, h, aler):
    sid = aler
    r = client.post(f"/stations/{sid}/revisions", json={"rev_no": "0", "date": "2026-08-20",
                    "description": "First issue"}, headers=h["dana"])
    assert r.status_code == 200 and r.json()["layout_version"] == "2.0.3"
    assert r.json()["created_by"] == "dana"
    assert client.post(f"/stations/{sid}/revisions", json={"rev_no": "0", "description": "dup"},
                       headers=h["dana"]).status_code == 409
    assert client.post(f"/stations/{sid}/revisions", json={"rev_no": "1", "date": "20/08/26",
                       "description": "x"}, headers=h["dana"]).status_code == 422
    assert client.post(f"/stations/{sid}/revisions", json={"rev_no": "1", "description": "x"},
                       headers=h["cara"]).status_code == 403
    rows = client.get(f"/stations/{sid}/revisions", headers=h["cara"]).json()
    assert [x["rev_no"] for x in rows] == ["0"]
    text = client.get(f"/stations/{sid}/reports/revisions", params={"format": "docx"},
                      headers=h["cara"]).content
    import docx
    body = "\n".join(c.text for t in docx.Document(io.BytesIO(text)).tables
                     for row in t.rows for c in row.cells)
    assert "First issue" in body
    # locked outside draft
    client.post(f"/stations/{sid}/self-check", json={}, headers=h["dana"])
    assert client.post(f"/stations/{sid}/revisions", json={"rev_no": "1", "description": "x"},
                       headers=h["dana"]).status_code == 409
    assert client.delete(f"/stations/{sid}/revisions/{rows[0]['id']}",
                         headers=h["dana"]).status_code == 409


def test_station_delete_cascades_reports(client, h, aler):
    client.post(f"/stations/{aler}/self-check", json={}, headers=h["dana"])
    assert list(client.upload_root.rglob("*.pdf"))
    assert client.delete(f"/stations/{aler}", headers=h["admin"]).status_code == 204
    assert not list(client.upload_root.rglob("*.pdf"))
